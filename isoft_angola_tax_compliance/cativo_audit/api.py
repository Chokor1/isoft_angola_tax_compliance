# Copyright (c) 2026, ISOFT and contributors
# For license information, please see license.txt
"""Server side of the Cativo Auditing page. Read-only."""

import frappe
from frappe import _

from isoft_angola_tax_compliance.cativo_audit import data, resolve

ROLES = ("System Manager", "Accounts Manager")


@frappe.whitelist()
def get_cativo_audit(company, from_date=None, to_date=None):
	"""Invoices with cativo in the period, and the cativo no invoice could be found for."""
	frappe.only_for(ROLES)
	if not company:
		frappe.throw(_("Choose a company."))

	accounts = data.get_accounts(company)
	warnings = []
	if accounts["derived"]:
		warnings.append(_("No saved accounts for {0}; using the company's setup.").format(company))
	if not accounts["cativo"]:
		warnings.append(_("No cativo account found for {0}. Add one in Cativo Audit Settings.").format(company))

	moves = data.load_movements(company, accounts)
	names = sorted({m["invoice"] for m in moves if m["invoice"]})
	invoices = data.load_invoices(names, accounts)
	return {
		"rows": resolve.build_rows(moves, invoices, from_date, to_date),
		"unassigned": resolve.unassigned(moves, from_date, to_date),
		"accounts": accounts,
		"warnings": warnings,
	}


@frappe.whitelist()
def export_cativo_audit(payload):
	"""Lay out in Excel exactly the table the page shows (filtered, sorted, with totals)."""
	frappe.only_for(ROLES)
	import json
	from io import BytesIO

	from openpyxl import Workbook
	from openpyxl.styles import Font

	data_ = json.loads(payload) if isinstance(payload, str) else payload
	columns = data_.get("columns") or []

	wb = Workbook()
	ws = wb.active
	ws.title = "Invoices with cativo"

	def put(row, col, value, kind):
		cell = ws.cell(row=row, column=col)
		if kind == "money" and value not in (None, ""):
			cell.value = float(value)
			cell.number_format = "#,##0.00"
		elif kind == "pct" and value not in (None, ""):
			cell.value = float(value)
			cell.number_format = "0.0%"
		else:
			# Forced to text so a value such as "=SUM(1)" is never read as a formula
			cell.value = "" if value is None else str(value)
			cell.data_type = "s"
		return cell

	for c, col in enumerate(columns, start=1):
		put(1, c, col.get("label"), "text").font = Font(bold=True)

	r = 1
	for row in data_.get("rows") or []:
		r += 1
		for c, col in enumerate(columns, start=1):
			put(r, c, row.get(col["key"]), col.get("kind"))

	totals = data_.get("totals") or {}
	r += 1
	for c, col in enumerate(columns, start=1):
		if col["key"] in totals:
			put(r, c, totals[col["key"]], col.get("kind")).font = Font(bold=True)

	unassigned = data_.get("unassigned") or {}
	r += 2
	put(r, 1, _("Cativo with no invoice identified"), "text").font = Font(bold=True)
	money_col = next((c for c, col in enumerate(columns, start=1) if col["key"] == "cativo"), 2)
	put(r, money_col, unassigned.get("amount") or 0, "money")
	if money_col != 2:
		put(r, 2, _("{0} entries").format(unassigned.get("count") or 0), "text")

	for c, col in enumerate(columns, start=1):
		ws.column_dimensions[ws.cell(row=1, column=c).column_letter].width = 14 if col.get("kind") in ("money", "pct") else 24

	out = BytesIO()
	wb.save(out)
	frappe.response["filename"] = "Cativo_Audit_%s_%s_%s.xlsx" % (
		data_.get("company") or "", data_.get("from_date") or "", data_.get("to_date") or ""
	)
	frappe.response["filecontent"] = out.getvalue()
	frappe.response["type"] = "binary"
